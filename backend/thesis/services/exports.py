"""Exportación editable de las revisiones existentes del proyecto."""

from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt

from thesis.models import RevisionCapitulo
from thesis.services.access import requerir_proyecto
from thesis.services.steps import PASOS_CAPITULO
from thesis.step_catalogue import CATALOGO_PASOS


def _agregar_markdown(document: Document, content: str) -> None:
    """Conserva el texto redactado; solo interpreta encabezados y listas simples."""
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=4)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)


def exportar_borrador_docx(actor, project_id) -> BytesIO:
    """Exporta el último texto guardado de cada paso como borrador, sin inventar contenido."""
    project = requerir_proyecto(actor, project_id)
    revisions = RevisionCapitulo.objects.filter(chapter__project=project).select_related("chapter").order_by(
        "chapter__number", "step_code", "-revision_no"
    )
    latest = {}
    for revision in revisions:
        latest.setdefault(revision.step_code, revision)

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(3)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 2
    document.add_heading(project.title or project.name, level=0)
    document.add_paragraph("Borrador de trabajo generado a partir de las revisiones guardadas.")
    document.add_paragraph(
        "Revise las fuentes, citas, contenido y formato exigido por su institución antes de presentar este documento."
    )
    for chapter_no, codes in PASOS_CAPITULO.items():
        present = [code for code in codes if code in latest]
        if not present:
            continue
        document.add_heading(f"Capítulo {chapter_no}", level=1)
        for code in present:
            document.add_heading(f"{code}. {CATALOGO_PASOS[code]['nombre']}", level=2)
            _agregar_markdown(document, latest[code].content_md)
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream
