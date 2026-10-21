"""Exportación editable de las revisiones existentes del proyecto."""

from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt

from thesis.models import RevisionCapitulo
from thesis.services.acceso import requerir_proyecto
from thesis.services.pasos import PASOS_CAPITULO
from thesis.catalogo_pasos import CATALOGO_PASOS


def _agregar_markdown(documento: Document, contenido: str) -> None:
    """Conserva el texto redactado; solo interpreta encabezados y listas simples."""
    for linea_original in contenido.splitlines():
        linea = linea_original.strip()
        if not linea:
            continue
        if linea.startswith("# "):
            documento.add_heading(linea[2:].strip(), level=3)
        elif linea.startswith("## "):
            documento.add_heading(linea[3:].strip(), level=4)
        elif linea.startswith(("- ", "* ")):
            documento.add_paragraph(linea[2:].strip(), style="List Bullet")
        else:
            documento.add_paragraph(linea)


def exportar_borrador_docx(usuario, id_proyecto) -> BytesIO:
    """Exporta el último texto guardado de cada paso como borrador, sin inventar contenido."""
    proyecto = requerir_proyecto(usuario, id_proyecto)
    revisiones = RevisionCapitulo.objects.filter(chapter__project=proyecto).select_related("chapter").order_by(
        "chapter__number", "step_code", "-revision_no"
    )
    ultimas = {}
    for revision in revisiones:
        ultimas.setdefault(revision.step_code, revision)

    documento = Document()
    seccion = documento.sections[0]
    seccion.page_width = Cm(21)
    seccion.page_height = Cm(29.7)
    seccion.left_margin = Cm(3)
    seccion.right_margin = Cm(2.5)
    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(3)
    for nombre_estilo in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet"):
        estilo = documento.styles[nombre_estilo]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(11)
        estilo.paragraph_format.line_spacing = 2
    documento.add_heading(proyecto.title or proyecto.name, level=0)
    documento.add_paragraph("Borrador de trabajo generado a partir de las revisiones guardadas.")
    documento.add_paragraph(
        "Revise las fuentes, citas, contenido y formato exigido por su institución antes de presentar este documento."
    )
    for numero_capitulo, codigos in PASOS_CAPITULO.items():
        presentes = [codigo for codigo in codigos if codigo in ultimas]
        if not presentes:
            continue
        documento.add_heading(f"Capítulo {numero_capitulo}", level=1)
        for codigo in presentes:
            documento.add_heading(f"{codigo}. {CATALOGO_PASOS[codigo]['nombre']}", level=2)
            _agregar_markdown(documento, ultimas[codigo].content_md)
    flujo = BytesIO()
    documento.save(flujo)
    flujo.seek(0)
    return flujo
